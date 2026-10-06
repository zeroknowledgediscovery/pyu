import pandas as pd, numpy as np, re, json, collections, warnings
from pathlib import Path
from openpyxl import load_workbook
warnings.filterwarnings('ignore')
B=Path('/mnt/data/joel_ua')
IN=B/'pyuria_longitudinal_full.csv'; OUT=B/'pyuria_longitudinal_corrected.csv'
F=pd.read_csv(IN,low_memory=False)
F['UA_DTM']=pd.to_datetime(F.UA_DTM,errors='coerce')
F['Anon_MRN']=pd.to_numeric(F.Anon_MRN,errors='coerce').astype('Int64')
F['Anon_Encounter_ID']=pd.to_numeric(F.Anon_Encounter_ID,errors='coerce').astype('Int64')
F['_row']=np.arange(len(F))
sp=pd.read_csv('/mnt/data/joel_mc/split_patients_seed20261006.csv')
train=set(pd.to_numeric(sp.loc[sp.split.astype(str).str.lower().eq('train'),'Anon_MRN'],errors='coerce').dropna().astype(int))
allm=set(F.Anon_MRN.dropna().astype(int))

def xlsx(path, cols):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]
    it=ws.iter_rows(values_only=True); hdr=list(next(it)); ids=[hdr.index(c) for c in cols]
    return pd.DataFrame((tuple(r[i] for i in ids) for r in it),columns=cols)

def safe(s): return re.sub(r'[^A-Za-z0-9]+','_',str(s)).strip('_')[:28]
# per-patient target info
ua={}
for m,g in F.dropna(subset=['Anon_MRN','UA_DTM']).groupby('Anon_MRN',sort=False):
    ua[int(m)] = (g['_row'].to_numpy(int), g['UA_DTM'].to_numpy('datetime64[ns]'), g['Anon_Encounter_ID'].astype('Int64').to_numpy())

def add_event_counts(prefix, h, timecol, windows=(30,365,None), catcol=None, cats=None, encounter_col=None, exclude_same_enc=False, safety_hours=0):
    h=h.dropna(subset=['Anon_MRN',timecol]).copy(); h['Anon_MRN']=pd.to_numeric(h.Anon_MRN,errors='coerce').astype('Int64'); h=h[h.Anon_MRN.isin(allm)]
    if catcol is None: groups=[(None,h)]
    else: groups=[(c,h[h[catcol].eq(c)]) for c in cats]
    for cat,sub in groups:
        arrays={w:np.zeros(len(F),dtype=np.int32) for w in windows}
        for m,g in sub.groupby('Anon_MRN',sort=False):
            m=int(m)
            if m not in ua: continue
            ii, uts, encids = ua[m]
            gg=g.sort_values(timecol)
            ts=gg[timecol].to_numpy('datetime64[ns]')
            evenc = gg[encounter_col].astype('Int64').to_numpy() if (encounter_col and encounter_col in gg.columns) else None
            for q,(row,ut) in enumerate(zip(ii,uts)):
                cutoff=ut-np.timedelta64(int(safety_hours*3600),'s') if safety_hours else ut
                mask=ts < cutoff
                if exclude_same_enc and evenc is not None and pd.notna(encids[q]):
                    mask &= (evenc != encids[q])
                valid_ts=ts[mask]
                for w in windows:
                    if w is None: n=len(valid_ts)
                    else: n=int(np.sum(valid_ts >= cutoff-np.timedelta64(int(w),'D')))
                    arrays[w][row]=n
        for w,a in arrays.items():
            suf='life' if w is None else f'{w}d'
            nm=f'{prefix}_{suf}' if cat is None else f'{prefix}_{safe(cat)}_{suf}'
            F[nm]=a

# Recompute encounter counts excluding current encounter.
print('encounters',flush=True)
enc=xlsx(B/'anonymized_PHI_UA_ENCOUNTER_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','ADMT_DT'])
enc['ADMT_DT']=pd.to_datetime(enc.ADMT_DT,errors='coerce'); enc['Anon_Encounter_ID']=pd.to_numeric(enc.Anon_Encounter_ID,errors='coerce').astype('Int64')
add_event_counts('prior_enc',enc,'ADMT_DT',(7,30,365,None),encounter_col='Anon_Encounter_ID',exclude_same_enc=True)

# Diagnoses: exclude current encounter completely. Add broad + phenotype categories + top ICD3.
print('diagnoses',flush=True)
dx=xlsx(B/'anonymized_PHI_UA_DIAGNOSIS_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','ADMT_DT','DIAGNOSIS'])
dx['ADMT_DT']=pd.to_datetime(dx.ADMT_DT,errors='coerce'); dx['Anon_Encounter_ID']=pd.to_numeric(dx.Anon_Encounter_ID,errors='coerce').astype('Int64'); dx['Anon_MRN']=pd.to_numeric(dx.Anon_MRN,errors='coerce').astype('Int64')
dx['icd3']=dx.DIAGNOSIS.fillna('').astype(str).str.upper().str.replace('.','',regex=False).str[:3]
add_event_counts('dx',dx,'ADMT_DT',(30,365,None),encounter_col='Anon_Encounter_ID',exclude_same_enc=True)
for nm,pat in {'uti':r'^N39','pyelo':r'^N1[0-2]','urinary':r'^R3','renal':r'^N[0-2]','fever':r'^R50','abd':r'^R10'}.items():
    add_event_counts('dx_'+nm,dx[dx.icd3.str.match(pat,na=False)],'ADMT_DT',(30,365,None),encounter_col='Anon_Encounter_ID',exclude_same_enc=True)
topdx=dx[dx.Anon_MRN.isin(train)&dx.icd3.ne('')].icd3.value_counts().head(20).index.tolist()
add_event_counts('dx3',dx,'ADMT_DT',(365,None),catcol='icd3',cats=topdx,encounter_col='Anon_Encounter_ID',exclude_same_enc=True)

# Problem list: exclude current encounter; use entry/noted datetime.
print('problem',flush=True)
pr=xlsx(B/'anonymized_PHI_UA_PROBLEM_LIST_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','PROBLEM_ICD10_LIST','CHRONIC_YN','NOTED_DATE','DATE_OF_ENTRY'])
pr['event_dt']=pd.to_datetime(pr.DATE_OF_ENTRY,errors='coerce').fillna(pd.to_datetime(pr.NOTED_DATE,errors='coerce')); pr['Anon_Encounter_ID']=pd.to_numeric(pr.Anon_Encounter_ID,errors='coerce').astype('Int64'); pr['Anon_MRN']=pd.to_numeric(pr.Anon_MRN,errors='coerce').astype('Int64')
pr['icd3']=pr.PROBLEM_ICD10_LIST.fillna('').astype(str).str.split(',').str[0].str.upper().str.replace('.','',regex=False).str[:3]
add_event_counts('problem',pr,'event_dt',(365,None),encounter_col='Anon_Encounter_ID',exclude_same_enc=True)
add_event_counts('problem_chronic',pr[pr.CHRONIC_YN.astype(str).str.upper().eq('Y')],'event_dt',(365,None),encounter_col='Anon_Encounter_ID',exclude_same_enc=True)
toppr=pr[pr.Anon_MRN.isin(train)&pr.icd3.ne('')].icd3.value_counts().head(15).index.tolist()
add_event_counts('pr3',pr,'event_dt',(365,None),catcol='icd3',cats=toppr,encounter_col='Anon_Encounter_ID',exclude_same_enc=True)

# Outpatient RX: no encounter id -> conservative 24h gap. Add top medication tokens.
print('outpatient rx',flush=True)
rx=xlsx(B/'anonymized_PHI_UA_OUTPATIENT_RX_MRNREFORMAT.xlsx',['Anon_MRN','ORDER_START_DATE','DISPLAY_NAME','DESCRIPTION'])
rx['event_dt']=pd.to_datetime(rx.ORDER_START_DATE,errors='coerce'); rx['Anon_MRN']=pd.to_numeric(rx.Anon_MRN,errors='coerce').astype('Int64')
txt=(rx.DISPLAY_NAME.fillna('').astype(str)+' '+rx.DESCRIPTION.fillna('').astype(str)).str.lower()
rx['drug']=txt.str.extract(r'([a-z][a-z0-9-]{2,})',expand=False).fillna('')
abxpat=r'amoxic|ampic|penicill|cephal|cef[a-z]|sulfameth|trimeth|nitrofur|ciprof|levof|azith|doxy|clinda|vancom|gentamic|tobram|meropen|piperac|linezolid|metronid'
rx['abx']=txt.str.contains(abxpat,regex=True,na=False)
add_event_counts('rx',rx,'event_dt',(30,365,None),safety_hours=24)
add_event_counts('rx_abx',rx[rx.abx],'event_dt',(30,365,None),safety_hours=24)
toprx=rx[rx.Anon_MRN.isin(train)&rx.drug.ne('')].drug.value_counts().head(20).index.tolist()
add_event_counts('rxdrug',rx,'event_dt',(30,365,None),catcol='drug',cats=toprx,safety_hours=24)

# Add top administered meds from precise timestamps.
print('med admin categories',flush=True)
mp=B/'anonymized_EX9576_MEDICATION_ADMIN.csv'; cc=collections.Counter()
for ch in pd.read_csv(mp,usecols=['Anon_MRN','MED_ORDER_NAME'],chunksize=700000,low_memory=False):
    ch['Anon_MRN']=pd.to_numeric(ch.Anon_MRN,errors='coerce'); ch=ch[ch.Anon_MRN.isin(train)]
    tok=ch.MED_ORDER_NAME.fillna('').astype(str).str.lower().str.extract(r'([a-z][a-z0-9-]{2,})',expand=False).fillna(''); cc.update(tok[tok.ne('')].value_counts().to_dict())
topmed=[x for x,_ in cc.most_common(15)]; parts=[]
for ch in pd.read_csv(mp,usecols=['Anon_MRN','MED_ADMIN_DTTM','MED_ORDER_NAME'],chunksize=700000,low_memory=False):
    ch['Anon_MRN']=pd.to_numeric(ch.Anon_MRN,errors='coerce').astype('Int64'); ch=ch[ch.Anon_MRN.isin(allm)]; ch['event_dt']=pd.to_datetime(ch.MED_ADMIN_DTTM,errors='coerce')
    ch['drug']=ch.MED_ORDER_NAME.fillna('').astype(str).str.lower().str.extract(r'([a-z][a-z0-9-]{2,})',expand=False).fillna(''); parts.append(ch[['Anon_MRN','event_dt','drug']])
med=pd.concat(parts,ignore_index=True); add_event_counts('meddrug',med,'event_dt',(7,30),catcol='drug',cats=topmed); del med,parts

# Procedures: service date lacks time. Current table already +1 day for broad counts; add top CPT3 with +1 day.
print('procedure categories',flush=True)
pp=B/'anonymized_EX9576_PROCEDURE.csv'; cc=collections.Counter()
for ch in pd.read_csv(pp,usecols=['Anon_MRN','CPT_CODE'],chunksize=700000,low_memory=False):
    ch['Anon_MRN']=pd.to_numeric(ch.Anon_MRN,errors='coerce'); ch=ch[ch.Anon_MRN.isin(train)]; tok=ch.CPT_CODE.fillna('').astype(str).str[:3]; cc.update(tok[tok.ne('')].value_counts().to_dict())
topproc=[x for x,_ in cc.most_common(15)]; parts=[]
for ch in pd.read_csv(pp,usecols=['Anon_MRN','SERVICE_DATE','CPT_CODE'],chunksize=700000,low_memory=False):
    ch['Anon_MRN']=pd.to_numeric(ch.Anon_MRN,errors='coerce').astype('Int64'); ch=ch[ch.Anon_MRN.isin(allm)]; ch['safe_dt']=pd.to_datetime(ch.SERVICE_DATE,errors='coerce')+pd.Timedelta(days=1); ch['cpt3']=ch.CPT_CODE.fillna('').astype(str).str[:3]; parts.append(ch[['Anon_MRN','safe_dt','cpt3']])
proc=pd.concat(parts,ignore_index=True); add_event_counts('cpt3',proc,'safe_dt',(365,None),catcol='cpt3',cats=topproc); del proc,parts

# Remove obviously unsafe/or bookkeeping fields from final CSV only later in model script; keep identifiers for reproducibility.
F.drop(columns=['_row'],inplace=True)
F.to_csv(OUT,index=False)
meta={'shape':F.shape,'topdx':topdx,'toppr':toppr,'toprx':toprx,'topmed':topmed,'topproc':topproc,
      'timing_rules':{'diagnosis':'exclude current encounter','problem':'exclude current encounter','rx':'24h safety gap','encounter':'exclude current encounter','medadmin':'exact timestamp','procedure':'service date +1 day'}}
(B/'pyuria_longitudinal_corrected_metadata.json').write_text(json.dumps(meta,indent=2))
print('DONE',F.shape,OUT,flush=True)
