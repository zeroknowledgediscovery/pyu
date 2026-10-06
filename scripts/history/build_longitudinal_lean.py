import pandas as pd, numpy as np, re, json, collections, warnings, time
from pathlib import Path
from openpyxl import load_workbook
warnings.filterwarnings('ignore')
B=Path('/mnt/data/joel_ua'); OUT=B/'pyuria_longitudinal_full.csv'

def xlsx(path, usecols=None):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; it=ws.iter_rows(values_only=True); hdr=list(next(it)); idx=list(range(len(hdr))) if usecols is None else [hdr.index(c) for c in usecols]; cols=[hdr[i] for i in idx]; rows=[tuple(r[i] for i in idx) for r in it]; return pd.DataFrame(rows,columns=cols)
def safe(x): return re.sub(r'[^A-Za-z0-9]+','_',str(x)).strip('_')[:34]
base=pd.read_csv(B/'pyuria_lsm_dataset.csv',low_memory=False); base['UA_DTM']=pd.to_datetime(base.UA_DTM,errors='coerce'); base['Anon_MRN']=pd.to_numeric(base.Anon_MRN,errors='coerce').astype('Int64'); base['Anon_Encounter_ID']=pd.to_numeric(base.Anon_Encounter_ID,errors='coerce').astype('Int64'); F=base.copy()
sp=pd.read_csv('/mnt/data/joel_mc/split_patients_seed20261006.csv'); train=set(pd.to_numeric(sp.loc[sp.split.astype(str).str.lower().eq('train'),'Anon_MRN'],errors='coerce').dropna().astype(int)); allm=set(base.Anon_MRN.dropna().astype(int))
ua={int(m):(np.asarray(idx,dtype=int),F.loc[idx,'UA_DTM'].values.astype('datetime64[ns]')) for m,idx in F.groupby('Anon_MRN').groups.items() if pd.notna(m)}

def add_counts(prefix,h,timecol,windows=(30,365,None),catcol=None,cats=None):
    arrays={}
    keys=[None] if catcol is None else list(cats)
    for k in keys:
        for w in windows: arrays[(k,w)]=np.zeros(len(F),np.int32)
    h=h[h.Anon_MRN.notna() & h[timecol].notna()].copy(); h.Anon_MRN=h.Anon_MRN.astype(int); h=h[h.Anon_MRN.isin(allm)]
    for m,g in h.groupby('Anon_MRN',sort=False):
        if m not in ua: continue
        ii,uts=ua[m]
        if catcol is None:
            t=np.sort(g[timecol].values.astype('datetime64[ns]')); r=np.searchsorted(t,uts,'left')
            for w in windows: arrays[(None,w)][ii]=r if w is None else r-np.searchsorted(t,uts-np.timedelta64(int(w),'D'),'left')
        else:
            for k,gg in g.groupby(catcol,sort=False):
                if k not in keys: continue
                t=np.sort(gg[timecol].values.astype('datetime64[ns]')); r=np.searchsorted(t,uts,'left')
                for w in windows: arrays[(k,w)][ii]=r if w is None else r-np.searchsorted(t,uts-np.timedelta64(int(w),'D'),'left')
    for (k,w),v in arrays.items():
        suffix='life' if w is None else f'{w}d'; name=f'{prefix}_{suffix}' if k is None else f'{prefix}_{safe(k)}_{suffix}'; F[name]=v

def topcat(df,col,n):
    z=df[df.Anon_MRN.isin(train)][col].fillna('').astype(str); return z[z.ne('')].value_counts().head(n).index.tolist()

def latest_numeric(prefix,h,namecol,timecol,valcol,names):
    for name in names:
        vals=np.full(len(F),np.nan); rec=np.full(len(F),np.nan); z=h[h[namecol].eq(name)].dropna(subset=[timecol]).sort_values(['Anon_MRN',timecol])
        for m,g in z.groupby('Anon_MRN',sort=False):
            if int(m) not in ua: continue
            ii,uts=ua[int(m)]; ts=g[timecol].values.astype('datetime64[ns]'); vv=pd.to_numeric(g[valcol],errors='coerce').values; j=np.searchsorted(ts,uts,'left')-1; ok=j>=0; vals[ii[ok]]=vv[j[ok]]; rec[ii[ok]]=(uts[ok]-ts[j[ok]])/np.timedelta64(1,'h')
        F[prefix+'_'+safe(name)]=vals; F[prefix+'rec_'+safe(name)]=rec

print('enc',flush=True); enc=xlsx(B/'anonymized_PHI_UA_ENCOUNTER_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','ADMT_DT','ZIP_CD','GENDER_IDENTITY']); enc['ADMT_DT']=pd.to_datetime(enc.ADMT_DT,errors='coerce'); enc.Anon_MRN=pd.to_numeric(enc.Anon_MRN,errors='coerce').astype('Int64'); enc.Anon_Encounter_ID=pd.to_numeric(enc.Anon_Encounter_ID,errors='coerce').astype('Int64'); F=F.merge(enc[['Anon_Encounter_ID','ZIP_CD','GENDER_IDENTITY']].drop_duplicates('Anon_Encounter_ID'),on='Anon_Encounter_ID',how='left'); F['ZIP3']=F.ZIP_CD.astype(str).str.replace(r'\.0$','',regex=True).str[:3].where(F.ZIP_CD.notna(),''); add_counts('prior_enc',enc,'ADMT_DT',(7,30,365,None))
print('dx',flush=True); dx=xlsx(B/'anonymized_PHI_UA_DIAGNOSIS_MRNREFORMAT.xlsx',['Anon_MRN','ADMT_DT','DIAGNOSIS']); dx.ADMT_DT=pd.to_datetime(dx.ADMT_DT,errors='coerce'); dx.Anon_MRN=pd.to_numeric(dx.Anon_MRN,errors='coerce').astype('Int64'); dx=dx[dx.Anon_MRN.isin(allm)]; dx['icd3']=dx.DIAGNOSIS.fillna('').astype(str).str.upper().str.replace('.','',regex=False).str[:3]; add_counts('dx',dx,'ADMT_DT',(30,365,None)); topdx=[]
for nm,pat in {'uti':r'^N39','pyelo':r'^N1[0-2]','urinary':r'^R3','fever':r'^R50','abd':r'^R10','renal':r'^N[0-2]'}.items(): add_counts('dx_'+nm,dx[dx.icd3.str.match(pat,na=False)],'ADMT_DT',(30,365,None))
print('problem',flush=True); pr=xlsx(B/'anonymized_PHI_UA_PROBLEM_LIST_MRNREFORMAT.xlsx',['Anon_MRN','PROBLEM_ICD10_LIST','CHRONIC_YN','NOTED_DATE','DATE_OF_ENTRY']); pr['event_dt']=pd.to_datetime(pr.DATE_OF_ENTRY,errors='coerce').fillna(pd.to_datetime(pr.NOTED_DATE,errors='coerce')); pr.Anon_MRN=pd.to_numeric(pr.Anon_MRN,errors='coerce').astype('Int64'); pr=pr[pr.Anon_MRN.isin(allm)]; pr['icd3']=pr.PROBLEM_ICD10_LIST.fillna('').astype(str).str.split(',').str[0].str.upper().str.replace('.','',regex=False).str[:3]; add_counts('problem',pr,'event_dt',(365,None)); add_counts('problem_chronic',pr[pr.CHRONIC_YN.astype(str).str.upper().eq('Y')],'event_dt',(365,None)); toppr=[]
print('rx',flush=True); rx=xlsx(B/'anonymized_PHI_UA_OUTPATIENT_RX_MRNREFORMAT.xlsx',['Anon_MRN','ORDER_START_DATE','ORDER_STOP_DATE','DISPLAY_NAME','DESCRIPTION']); rx['event_dt']=pd.to_datetime(rx.ORDER_START_DATE,errors='coerce'); rx['stop_dt']=pd.to_datetime(rx.ORDER_STOP_DATE,errors='coerce'); rx.Anon_MRN=pd.to_numeric(rx.Anon_MRN,errors='coerce').astype('Int64'); rx=rx[rx.Anon_MRN.isin(allm)]; rn=(rx.DISPLAY_NAME.fillna('').astype(str)+' '+rx.DESCRIPTION.fillna('').astype(str)).str.lower(); rx['drug']=rn.str.extract(r'([a-z][a-z0-9-]{2,})',expand=False).fillna(''); abx=r'amoxic|ampic|penicill|cephal|cef[a-z]|sulfameth|trimeth|nitrofur|ciprof|levof|azith|doxy|clinda|vancom|gentamic|tobram|meropen|piperac|linezolid|metronid'; rx['abx']=rn.str.contains(abx,regex=True,na=False); add_counts('rx',rx,'event_dt',(30,365,None)); add_counts('rx_abx',rx[rx.abx],'event_dt',(30,365,None)); toprx=[]
print('img',flush=True); im=xlsx(B/'anonymized_PHI_UA_IMAGING_EPIC_MRNREFORMAT.xlsx',['Anon_MRN','FINAL_DT_TM','PROCDR1_NAME']); im.FINAL_DT_TM=pd.to_datetime(im.FINAL_DT_TM,errors='coerce'); im.Anon_MRN=pd.to_numeric(im.Anon_MRN,errors='coerce').astype('Int64'); im=im[im.Anon_MRN.isin(allm)]; im['mod']=im.PROCDR1_NAME.fillna('').astype(str).str.upper().str.extract(r'^(XR|CT|US|MRI|MR |ECHO|NM|FLUOR)',expand=False).fillna('OTHER').str.strip(); add_counts('img',im,'FINAL_DT_TM',(30,365,None))
print('med',flush=True); mp=B/'anonymized_EX9576_MEDICATION_ADMIN.csv'; cc=collections.Counter()
for ch in pd.read_csv(mp,usecols=['Anon_MRN','MED_ORDER_NAME'],chunksize=600000,low_memory=False): ch.Anon_MRN=pd.to_numeric(ch.Anon_MRN,errors='coerce'); ch=ch[ch.Anon_MRN.isin(train)]; tok=ch.MED_ORDER_NAME.fillna('').astype(str).str.lower().str.extract(r'([a-z][a-z0-9-]{2,})',expand=False).fillna(''); cc.update(tok[tok.ne('')].value_counts().to_dict())
topmed=[x for x,_ in cc.most_common(12)]; parts=[]
for ch in pd.read_csv(mp,usecols=['Anon_MRN','MED_ADMIN_DTTM','MED_ORDER_NAME'],chunksize=600000,low_memory=False): ch.Anon_MRN=pd.to_numeric(ch.Anon_MRN,errors='coerce').astype('Int64'); ch=ch[ch.Anon_MRN.isin(allm)]; ch['event_dt']=pd.to_datetime(ch.MED_ADMIN_DTTM,errors='coerce'); nm=ch.MED_ORDER_NAME.fillna('').astype(str).str.lower(); ch['drug']=nm.str.extract(r'([a-z][a-z0-9-]{2,})',expand=False).fillna(''); ch['abx']=nm.str.contains(abx,regex=True,na=False); parts.append(ch[['Anon_MRN','event_dt','drug','abx']])
med=pd.concat(parts,ignore_index=True); add_counts('med',med,'event_dt',(1,7,30,365,None)); add_counts('med_abx',med[med.abx],'event_dt',(1,7,30,365,None)); del med,parts
print('proc',flush=True); pp=B/'anonymized_EX9576_PROCEDURE.csv'; cc=collections.Counter()
for ch in pd.read_csv(pp,usecols=['Anon_MRN','CPT_CODE'],chunksize=600000,low_memory=False): ch.Anon_MRN=pd.to_numeric(ch.Anon_MRN,errors='coerce'); ch=ch[ch.Anon_MRN.isin(train)]; p=ch.CPT_CODE.fillna('').astype(str).str[:3]; cc.update(p[p.ne('')].value_counts().to_dict())
topproc=[x for x,_ in cc.most_common(12)]; parts=[]
for ch in pd.read_csv(pp,usecols=['Anon_MRN','SERVICE_DATE','CPT_CODE'],chunksize=600000,low_memory=False): ch.Anon_MRN=pd.to_numeric(ch.Anon_MRN,errors='coerce').astype('Int64'); ch=ch[ch.Anon_MRN.isin(allm)]; ch['safe_dt']=pd.to_datetime(ch.SERVICE_DATE,errors='coerce')+pd.Timedelta(days=1); ch['cpt3']=ch.CPT_CODE.fillna('').astype(str).str[:3]; parts.append(ch[['Anon_MRN','safe_dt','cpt3']])
proc=pd.concat(parts,ignore_index=True); add_counts('proc',proc,'safe_dt',(30,365,None)); del proc,parts
print('labs pass1',flush=True); lp=B/'anonymized_EX9576_LABS.csv'; cc=collections.Counter()
for ch in pd.read_csv(lp,usecols=['Anon_MRN','LAB_NAME','VALUE_NUM'],chunksize=700000,low_memory=False): ch.Anon_MRN=pd.to_numeric(ch.Anon_MRN,errors='coerce'); ch=ch[ch.Anon_MRN.isin(train)&ch.VALUE_NUM.notna()]; n=ch.LAB_NAME.fillna('').astype(str); n=n[~n.str.contains('urine',case=False,na=False)]; cc.update(n.value_counts().to_dict())
toplab=[x for x,_ in cc.most_common(20)]; lpv=[]; up=[]
print('labs pass2',flush=True)
for ch in pd.read_csv(lp,usecols=['Anon_MRN','ORDR_PERFRMD_DT_TM','LAB_NAME','VALUE_NUM','VALUE_TXT'],chunksize=700000,low_memory=False):
 ch.Anon_MRN=pd.to_numeric(ch.Anon_MRN,errors='coerce').astype('Int64')
 ch=ch[ch.Anon_MRN.isin(allm)]
 ch['event_dt']=pd.to_datetime(ch.ORDR_PERFRMD_DT_TM,errors='coerce')
 n=ch.LAB_NAME.fillna('').astype(str)
 s=n.isin(toplab)&ch.VALUE_NUM.notna()
 u=n.isin(['WBC, Urine','External Urine WBC'])
 if s.any(): lpv.append(ch.loc[s,['Anon_MRN','event_dt','LAB_NAME','VALUE_NUM']])
 if u.any(): up.append(ch.loc[u,['Anon_MRN','event_dt','VALUE_NUM','VALUE_TXT']])
labs=pd.concat(lpv,ignore_index=True); add_counts('labtop',labs,'event_dt',(1,7,30,365,None)); latest_numeric('lab',labs,'LAB_NAME','event_dt','VALUE_NUM',toplab)
ur=pd.concat(up,ignore_index=True)
def uw(a,b):
 if pd.notna(a): return float(a)
 s=str(b).lower().replace(' ',''); m=re.search(r'(\d+)[-–](\d+)',s)
 if m:return (float(m.group(1))+float(m.group(2)))/2
 m=re.search(r'>=?([0-9]+)',s)
 if m:return float(m.group(1))+1
 m=re.search(r'([0-9]+)',s); return float(m.group(1)) if m else np.nan
ur['w']=[uw(a,b) for a,b in zip(ur.VALUE_NUM,ur.VALUE_TXT)]; cnt=np.zeros(len(F),np.int32); pos=np.zeros(len(F),np.int32); last=np.full(len(F),np.nan); rec=np.full(len(F),np.nan)
for m,g in ur.dropna(subset=['event_dt']).sort_values(['Anon_MRN','event_dt']).groupby('Anon_MRN',sort=False):
 if int(m) not in ua: continue
 ii,uts=ua[int(m)]; ts=g.event_dt.values.astype('datetime64[ns]'); vv=np.asarray(g.w,float); cuts=uts-np.timedelta64(24,'h'); jj=np.searchsorted(ts,cuts,'left')
 for q,i in enumerate(ii):
  j=jj[q]; x=vv[:j]; cnt[i]=np.isfinite(x).sum(); pos[i]=np.sum(x[np.isfinite(x)]>5); k=j-1
  while k>=0 and not np.isfinite(vv[k]): k-=1
  if k>=0: last[i]=vv[k]; rec[i]=(uts[q]-ts[k])/np.timedelta64(1,'D')
F['prior_ua_count']=cnt; F['prior_pyuria_count']=pos; F['prior_pyuria_rate']=np.divide(pos,cnt,out=np.zeros(len(F),float),where=cnt>0); F['prior_ua_wbc_last']=last; F['prior_ua_rec_days']=rec
F=F.replace([np.inf,-np.inf],np.nan); F.to_csv(OUT,index=False); meta={'shape':F.shape,'topdx':topdx,'toppr':toppr,'toprx':toprx,'topmed':topmed,'topproc':topproc,'toplab':toplab}; (B/'pyuria_longitudinal_metadata.json').write_text(json.dumps(meta,indent=2)); print('DONE',F.shape,flush=True)
