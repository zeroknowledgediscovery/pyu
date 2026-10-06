import pandas as pd, numpy as np, re, json, collections, warnings
from pathlib import Path
from openpyxl import load_workbook
warnings.filterwarnings('ignore')
B=Path('/mnt/data/joel_ua'); OUT=B/'pyuria_longitudinal_full.csv'

def xlsx(path, cols):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; it=ws.iter_rows(values_only=True); hdr=list(next(it)); ids=[hdr.index(c) for c in cols]
    df=pd.DataFrame((tuple(r[i] for i in ids) for r in it),columns=cols); wb.close(); return df
def dt(x): return pd.to_datetime(x,errors='coerce',format='mixed')
def mrn(x): return pd.to_numeric(x,errors='coerce').astype('Int64')
def safe(x): return re.sub(r'[^A-Za-z0-9]+','_',str(x)).strip('_')[:34]
F=pd.read_csv(B/'pyuria_lsm_dataset.csv',low_memory=False); F['UA_DTM']=dt(F.UA_DTM); F['Anon_MRN']=mrn(F.Anon_MRN); F['Anon_Encounter_ID']=mrn(F.Anon_Encounter_ID); allm=set(F.Anon_MRN.dropna().astype(int))
sp=pd.read_csv('/mnt/data/joel_mc/split_patients_seed20261006.csv'); train=set(pd.to_numeric(sp.loc[sp.split.astype(str).str.lower().eq('train'),'Anon_MRN'],errors='coerce').dropna().astype(int))
ua={int(m):(np.asarray(idx,dtype=int),F.loc[idx,'UA_DTM'].values.astype('datetime64[ns]')) for m,idx in F.groupby('Anon_MRN').groups.items() if pd.notna(m)}

def count_windows(t,uts,windows):
    t=np.sort(np.asarray(t,dtype='datetime64[ns]')); r=np.searchsorted(t,uts,'left'); out={}
    for w in windows:
        out[w]=r if w is None else r-np.searchsorted(t,uts-np.timedelta64(int(w),'D'),'left')
    return out

def add_source_counts(h,timecol,prefix,windows=(30,365,None),mask_defs=None):
    h=h[h.Anon_MRN.notna() & h[timecol].notna()].copy(); h.Anon_MRN=h.Anon_MRN.astype(int); h=h[h.Anon_MRN.isin(allm)]
    def fill(z,nm):
        arrays={w:np.zeros(len(F),np.int32) for w in windows}
        for m,g in z.groupby('Anon_MRN',sort=False):
            if m not in ua: continue
            ii,uts=ua[m]; c=count_windows(g[timecol].values.astype('datetime64[ns]'),uts,windows)
            for w,v in c.items(): arrays[w][ii]=v
        for w,v in arrays.items(): F[nm+('_life' if w is None else f'_{w}d')]=v
    fill(h,prefix)
    if mask_defs:
        for nm,func in mask_defs.items():
            mask=np.asarray(func(h),dtype=bool)
            fill(h.loc[mask],prefix+'_'+nm)

print('enc',flush=True)
enc=xlsx(B/'anonymized_PHI_UA_ENCOUNTER_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','ADMT_DT','ZIP_CD','GENDER_IDENTITY']); enc['Anon_MRN']=mrn(enc.Anon_MRN); enc['Anon_Encounter_ID']=mrn(enc.Anon_Encounter_ID); enc['ADMT_DT']=dt(enc.ADMT_DT); F=F.merge(enc[['Anon_Encounter_ID','ZIP_CD','GENDER_IDENTITY']].drop_duplicates('Anon_Encounter_ID'),on='Anon_Encounter_ID',how='left'); F['ZIP3']=F.ZIP_CD.astype(str).str.replace(r'\.0$','',regex=True).str[:3].where(F.ZIP_CD.notna(),''); add_source_counts(enc,'ADMT_DT','prior_enc',(7,30,365,None))

print('dx',flush=True)
dx=xlsx(B/'anonymized_PHI_UA_DIAGNOSIS_MRNREFORMAT.xlsx',['Anon_MRN','ADMT_DT','DIAGNOSIS']); dx['Anon_MRN']=mrn(dx.Anon_MRN); dx['ADMT_DT']=dt(dx.ADMT_DT); dx['icd']=dx.DIAGNOSIS.fillna('').astype(str).str.upper().str.replace('.','',regex=False)
dx['uti']=dx.icd.str.match(r'^N39',na=False); dx['pyelo']=dx.icd.str.match(r'^N1[0-2]',na=False); dx['urinary']=dx.icd.str.match(r'^R3',na=False); dx['renal']=dx.icd.str.match(r'^N[0-2]',na=False)
add_source_counts(dx,'ADMT_DT','dx',(30,365,None),{'uti':lambda g:g.uti,'pyelo':lambda g:g.pyelo,'urinary':lambda g:g.urinary,'renal':lambda g:g.renal}); del dx

print('problem',flush=True)
pr=xlsx(B/'anonymized_PHI_UA_PROBLEM_LIST_MRNREFORMAT.xlsx',['Anon_MRN','CHRONIC_YN','NOTED_DATE','DATE_OF_ENTRY']); pr['Anon_MRN']=mrn(pr.Anon_MRN); pr['event_dt']=dt(pr.DATE_OF_ENTRY).fillna(dt(pr.NOTED_DATE)); add_source_counts(pr,'event_dt','problem',(365,None),{'chronic':lambda g:g.CHRONIC_YN.astype(str).str.upper().eq('Y')}); del pr

print('rx',flush=True)
rx=xlsx(B/'anonymized_PHI_UA_OUTPATIENT_RX_MRNREFORMAT.xlsx',['Anon_MRN','ORDER_START_DATE','DISPLAY_NAME','DESCRIPTION']); rx['Anon_MRN']=mrn(rx.Anon_MRN); rx['event_dt']=dt(rx.ORDER_START_DATE); txt=(rx.DISPLAY_NAME.fillna('').astype(str)+' '+rx.DESCRIPTION.fillna('').astype(str)).str.lower(); abxpat=r'amoxic|ampic|penicill|cephal|cef[a-z]|sulfameth|trimeth|nitrofur|ciprof|levof|azith|doxy|clinda|vancom|gentamic|tobram|meropen|piperac|linezolid|metronid'; rx['abx']=txt.str.contains(abxpat,regex=True,na=False); add_source_counts(rx,'event_dt','rx',(30,365,None),{'abx':lambda g:g.abx}); del rx

print('imaging',flush=True)
im=xlsx(B/'anonymized_PHI_UA_IMAGING_EPIC_MRNREFORMAT.xlsx',['Anon_MRN','FINAL_DT_TM','PROCDR1_NAME']); im['Anon_MRN']=mrn(im.Anon_MRN); im['event_dt']=dt(im.FINAL_DT_TM); im['name']=im.PROCDR1_NAME.fillna('').astype(str).str.upper(); add_source_counts(im,'event_dt','imaging',(30,365,None),{'us':lambda g:g.name.str.startswith('US'),'ct':lambda g:g.name.str.startswith('CT'),'xr':lambda g:g.name.str.startswith('XR'),'mri':lambda g:g.name.str.match(r'^(MRI|MR )',na=False)}); del im

print('medadmin',flush=True)
parts=[]
for ch in pd.read_csv(B/'anonymized_EX9576_MEDICATION_ADMIN.csv',usecols=['Anon_MRN','MED_ADMIN_DTTM','MED_ORDER_NAME'],chunksize=800000,low_memory=False):
    ch['Anon_MRN']=mrn(ch.Anon_MRN); ch=ch[ch.Anon_MRN.isin(allm)]; ch['event_dt']=dt(ch.MED_ADMIN_DTTM); ch['abx']=ch.MED_ORDER_NAME.fillna('').astype(str).str.lower().str.contains(abxpat,regex=True,na=False); parts.append(ch[['Anon_MRN','event_dt','abx']])
med=pd.concat(parts,ignore_index=True); add_source_counts(med,'event_dt','medadmin',(1,7,30,365,None),{'abx':lambda g:g.abx}); del med,parts

print('procedure',flush=True)
parts=[]
for ch in pd.read_csv(B/'anonymized_EX9576_PROCEDURE.csv',usecols=['Anon_MRN','SERVICE_DATE'],chunksize=800000,low_memory=False):
    ch['Anon_MRN']=mrn(ch.Anon_MRN); ch=ch[ch.Anon_MRN.isin(allm)]; ch['event_dt']=dt(ch.SERVICE_DATE)+pd.Timedelta(days=1); parts.append(ch[['Anon_MRN','event_dt']])
proc=pd.concat(parts,ignore_index=True); add_source_counts(proc,'event_dt','procedure',(30,365,None)); del proc,parts

print('labs1',flush=True)
cc=collections.Counter(); lp=B/'anonymized_EX9576_LABS.csv'
for ch in pd.read_csv(lp,usecols=['Anon_MRN','LAB_NAME','VALUE_NUM'],chunksize=900000,low_memory=False):
    ch['Anon_MRN']=pd.to_numeric(ch.Anon_MRN,errors='coerce'); ch=ch[ch.Anon_MRN.isin(train)&ch.VALUE_NUM.notna()]; n=ch.LAB_NAME.fillna('').astype(str); n=n[~n.str.contains('urine',case=False,na=False)]; cc.update(n.value_counts().to_dict())
toplab=[x for x,_ in cc.most_common(20)]; parts=[]; urparts=[]; countparts=[]
print('labs2',flush=True)
for ch in pd.read_csv(lp,usecols=['Anon_MRN','ORDR_PERFRMD_DT_TM','LAB_NAME','VALUE_NUM','VALUE_TXT'],chunksize=900000,low_memory=False):
    ch['Anon_MRN']=mrn(ch.Anon_MRN); ch=ch[ch.Anon_MRN.isin(allm)]; ch['event_dt']=dt(ch.ORDR_PERFRMD_DT_TM); countparts.append(ch[['Anon_MRN','event_dt']]); n=ch.LAB_NAME.fillna('').astype(str); s=n.isin(toplab)&ch.VALUE_NUM.notna(); u=n.isin(['WBC, Urine','External Urine WBC']);
    if s.any(): parts.append(ch.loc[s,['Anon_MRN','event_dt','LAB_NAME','VALUE_NUM']]);
    if u.any(): urparts.append(ch.loc[u,['Anon_MRN','event_dt','VALUE_NUM','VALUE_TXT']])
alllabs=pd.concat(countparts,ignore_index=True); add_source_counts(alllabs,'event_dt','lab',(1,7,30,365,None)); labs=pd.concat(parts,ignore_index=True)
# latest top lab values per patient, all 20 at once
for m,g in labs.dropna(subset=['event_dt']).sort_values(['Anon_MRN','event_dt']).groupby('Anon_MRN',sort=False):
    if int(m) not in ua: continue
    ii,uts=ua[int(m)]
    for lname,gg in g.groupby('LAB_NAME',sort=False):
        if lname not in toplab: continue
        ts=gg.event_dt.values.astype('datetime64[ns]'); vv=pd.to_numeric(gg.VALUE_NUM,errors='coerce').values; j=np.searchsorted(ts,uts,'left')-1; ok=j>=0; col='lab_'+safe(lname); rec='labrec_'+safe(lname)
        if col not in F: F[col]=np.nan; F[rec]=np.nan
        F.loc[ii[ok],col]=vv[j[ok]]; F.loc[ii[ok],rec]=(uts[ok]-ts[j[ok]])/np.timedelta64(1,'h')
# prior urine WBC >=24h old
ur=pd.concat(urparts,ignore_index=True)
def pw(a,b):
    if pd.notna(a): return float(a)
    s=str(b).lower().replace(' ',''); m=re.search(r'(\d+)[-–](\d+)',s)
    if m:return (float(m.group(1))+float(m.group(2)))/2
    m=re.search(r'>=?([0-9]+)',s)
    if m:return float(m.group(1))+1
    m=re.search(r'([0-9]+)',s); return float(m.group(1)) if m else np.nan
ur['wbc_est']=[pw(a,b) for a,b in zip(ur.VALUE_NUM,ur.VALUE_TXT)]; cnt=np.zeros(len(F),np.int32); pos=np.zeros(len(F),np.int32); last=np.full(len(F),np.nan); rec=np.full(len(F),np.nan)
for m,g in ur.dropna(subset=['event_dt']).sort_values(['Anon_MRN','event_dt']).groupby('Anon_MRN',sort=False):
    if int(m) not in ua: continue
    ii,uts=ua[int(m)]; ts=g.event_dt.values.astype('datetime64[ns]'); vv=np.asarray(g.wbc_est,float); jj=np.searchsorted(ts,uts-np.timedelta64(24,'h'),'left')
    for q,i in enumerate(ii):
        j=jj[q]; x=vv[:j]; cnt[i]=np.isfinite(x).sum(); pos[i]=np.sum(x[np.isfinite(x)]>5); k=j-1
        while k>=0 and not np.isfinite(vv[k]): k-=1
        if k>=0: last[i]=vv[k]; rec[i]=(uts[q]-ts[k])/np.timedelta64(1,'D')
F['prior_ua_count']=cnt; F['prior_pyuria_count']=pos; F['prior_pyuria_rate']=np.divide(pos,cnt,out=np.zeros(len(F),float),where=cnt>0); F['prior_ua_wbc_last']=last; F['prior_ua_rec_days']=rec
F=F.replace([np.inf,-np.inf],np.nan); F.to_csv(OUT,index=False); (B/'pyuria_longitudinal_metadata.json').write_text(json.dumps({'shape':list(F.shape),'top_labs':toplab},indent=2)); print('DONE',F.shape,toplab,flush=True)
