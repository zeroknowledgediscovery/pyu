import pandas as pd, numpy as np, re, json, warnings
from pathlib import Path
from openpyxl import load_workbook
warnings.filterwarnings('ignore')
B=Path('/mnt/data/joel_ua'); OUT=B/'pyuria_longitudinal_corrected.csv'
F=pd.read_csv(B/'pyuria_longitudinal_full.csv',low_memory=False)
F['UA_DTM']=pd.to_datetime(F.UA_DTM,errors='coerce'); F['Anon_MRN']=pd.to_numeric(F.Anon_MRN,errors='coerce').astype('Int64'); F['Anon_Encounter_ID']=pd.to_numeric(F.Anon_Encounter_ID,errors='coerce').astype('Int64')
F['_row']=np.arange(len(F))

def xlsx(path, cols):
 wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; it=ws.iter_rows(values_only=True); hdr=list(next(it)); ids=[hdr.index(c) for c in cols]
 return pd.DataFrame((tuple(r[i] for i in ids) for r in it),columns=cols)

targets={}
for m,g in F.dropna(subset=['Anon_MRN','UA_DTM']).groupby('Anon_MRN',sort=False):
 targets[int(m)] = (g['_row'].to_numpy(int), g.UA_DTM.to_numpy('datetime64[ns]'), g.Anon_Encounter_ID.astype('Int64').to_numpy())

def count_source(df,timecol,prefix,windows=(30,365,None), enc_col=None, exclude_same=False, shift_hours=0, filters=None):
 # filters dict suffix -> boolean array/series on df; None gives base
 if filters is None: filters={'':np.ones(len(df),dtype=bool)}
 for suffix,mask in filters.items():
  sub=df.loc[np.asarray(mask)].dropna(subset=['Anon_MRN',timecol]).copy()
  groups={int(m):g for m,g in sub.groupby('Anon_MRN',sort=False)}
  arrs={w:np.zeros(len(F),dtype=np.int32) for w in windows}
  for m,(rows,uts,encs) in targets.items():
   g=groups.get(m)
   if g is None: continue
   ts=g[timecol].to_numpy('datetime64[ns]'); order=np.argsort(ts); ts=ts[order]
   evenc=None
   if enc_col:
    evenc=g[enc_col].astype('Int64').to_numpy()[order]
   for ridx,ut,eid in zip(rows,uts,encs):
    cutoff=ut-np.timedelta64(int(shift_hours*3600),'s')
    hi=np.searchsorted(ts,cutoff,side='left')
    for w in windows:
     lo=0 if w is None else np.searchsorted(ts,cutoff-np.timedelta64(int(w),'D'),side='left')
     n=hi-lo
     if exclude_same and evenc is not None and pd.notna(eid) and n>0:
      n-=int(np.sum(evenc[lo:hi]==eid))
     arrs[w][ridx]=max(0,n)
  stem=prefix+suffix
  for w,a in arrs.items():
   F[stem+('_life' if w is None else f'_{w}d')]=a

print('enc',flush=True)
enc=xlsx(B/'anonymized_PHI_UA_ENCOUNTER_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','ADMT_DT']); enc['Anon_MRN']=pd.to_numeric(enc.Anon_MRN,errors='coerce').astype('Int64'); enc['Anon_Encounter_ID']=pd.to_numeric(enc.Anon_Encounter_ID,errors='coerce').astype('Int64'); enc['event_dt']=pd.to_datetime(enc.ADMT_DT,errors='coerce')
count_source(enc,'event_dt','prior_enc',(7,30,365,None),'Anon_Encounter_ID',True)

print('dx',flush=True)
dx=xlsx(B/'anonymized_PHI_UA_DIAGNOSIS_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','ADMT_DT','DIAGNOSIS']); dx['Anon_MRN']=pd.to_numeric(dx.Anon_MRN,errors='coerce').astype('Int64'); dx['Anon_Encounter_ID']=pd.to_numeric(dx.Anon_Encounter_ID,errors='coerce').astype('Int64'); dx['event_dt']=pd.to_datetime(dx.ADMT_DT,errors='coerce'); icd=dx.DIAGNOSIS.fillna('').astype(str).str.upper().str.replace('.','',regex=False)
filters={'':np.ones(len(dx),bool),'_uti':icd.str.match(r'^N39',na=False),'_pyelo':icd.str.match(r'^N1[0-2]',na=False),'_urinary':icd.str.match(r'^R3',na=False),'_renal':icd.str.match(r'^N[0-2]',na=False),'_fever':icd.str.match(r'^R50',na=False),'_abd':icd.str.match(r'^R10',na=False)}
count_source(dx,'event_dt','dx',(30,365,None),'Anon_Encounter_ID',True,filters=filters)

print('problem',flush=True)
pr=xlsx(B/'anonymized_PHI_UA_PROBLEM_LIST_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','CHRONIC_YN','NOTED_DATE','DATE_OF_ENTRY']); pr['Anon_MRN']=pd.to_numeric(pr.Anon_MRN,errors='coerce').astype('Int64'); pr['Anon_Encounter_ID']=pd.to_numeric(pr.Anon_Encounter_ID,errors='coerce').astype('Int64'); pr['event_dt']=pd.to_datetime(pr.DATE_OF_ENTRY,errors='coerce').fillna(pd.to_datetime(pr.NOTED_DATE,errors='coerce'))
count_source(pr,'event_dt','problem',(365,None),'Anon_Encounter_ID',True)
count_source(pr,'event_dt','problem_chronic',(365,None),'Anon_Encounter_ID',True,filters={'':pr.CHRONIC_YN.astype(str).str.upper().eq('Y')})

print('rx',flush=True)
rx=xlsx(B/'anonymized_PHI_UA_OUTPATIENT_RX_MRNREFORMAT.xlsx',['Anon_MRN','ORDER_START_DATE','DISPLAY_NAME','DESCRIPTION']); rx['Anon_MRN']=pd.to_numeric(rx.Anon_MRN,errors='coerce').astype('Int64'); rx['event_dt']=pd.to_datetime(rx.ORDER_START_DATE,errors='coerce'); txt=(rx.DISPLAY_NAME.fillna('').astype(str)+' '+rx.DESCRIPTION.fillna('').astype(str)).str.lower(); abx=r'amoxic|ampic|penicill|cephal|cef[a-z]|sulfameth|trimeth|nitrofur|ciprof|levof|azith|doxy|clinda|vancom|gentamic|tobram|meropen|piperac|linezolid|metronid'
count_source(rx,'event_dt','rx',(30,365,None),shift_hours=24)
count_source(rx,'event_dt','rx_abx',(30,365,None),shift_hours=24,filters={'':txt.str.contains(abx,regex=True,na=False)})

F.drop(columns=['_row'],inplace=True); F.to_csv(OUT,index=False)
meta={'shape':F.shape,'timing_rules':{'encounter':'exclude current encounter','diagnosis':'exclude current encounter','problem':'exclude current encounter','outpatient_rx':'24h safety gap','lab_medadmin_imaging':'existing exact pre-UA censoring','procedure':'existing service-date +1d safety'}}
(B/'pyuria_longitudinal_corrected_metadata.json').write_text(json.dumps(meta,indent=2)); print('DONE',F.shape,flush=True)
