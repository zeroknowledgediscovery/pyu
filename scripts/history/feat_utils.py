import pandas as pd, numpy as np, re
from pathlib import Path
from openpyxl import load_workbook
B=Path('/mnt/data/joel_ua')
def xlsx(path, cols):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; it=ws.iter_rows(values_only=True); hdr=list(next(it)); ids=[hdr.index(c) for c in cols]
    df=pd.DataFrame((tuple(r[i] for i in ids) for r in it),columns=cols); wb.close(); return df
def dt(x): return pd.to_datetime(x,errors='coerce',format='mixed')
def mrn(x): return pd.to_numeric(x,errors='coerce').astype('Int64')
def base_events():
    F=pd.read_csv(B/'pyuria_lsm_dataset.csv',usecols=['Anon_MRN','UA_DTM']); F['Anon_MRN']=mrn(F.Anon_MRN); F['UA_DTM']=dt(F.UA_DTM); F['_row']=np.arange(len(F)); ua={int(m):(np.asarray(idx,dtype=int),F.loc[idx,'UA_DTM'].values.astype('datetime64[ns]')) for m,idx in F.groupby('Anon_MRN').groups.items() if pd.notna(m)}; return F,ua,set(F.Anon_MRN.dropna().astype(int))
def count_windows(t,uts,windows):
    t=np.sort(np.asarray(t,dtype='datetime64[ns]')); r=np.searchsorted(t,uts,'left'); return {w:(r if w is None else r-np.searchsorted(t,uts-np.timedelta64(int(w),'D'),'left')) for w in windows}
def source_counts(h,timecol,prefix,windows,ua,nrows,allm,mask_cols=None):
    out={prefix+('_life' if w is None else f'_{w}d'):np.zeros(nrows,np.int32) for w in windows}
    if mask_cols:
        for nm in mask_cols:
            for w in windows: out[prefix+'_'+nm+('_life' if w is None else f'_{w}d')]=np.zeros(nrows,np.int32)
    h=h[h.Anon_MRN.notna() & h[timecol].notna()].copy(); h.Anon_MRN=h.Anon_MRN.astype(int); h=h[h.Anon_MRN.isin(allm)]
    for m,g in h.groupby('Anon_MRN',sort=False):
        if m not in ua: continue
        ii,uts=ua[m]; c=count_windows(g[timecol].values.astype('datetime64[ns]'),uts,windows)
        for w,v in c.items(): out[prefix+('_life' if w is None else f'_{w}d')][ii]=v
        if mask_cols:
            for nm in mask_cols:
                gg=g[g[nm].fillna(False).astype(bool)]
                if gg.empty: continue
                c=count_windows(gg[timecol].values.astype('datetime64[ns]'),uts,windows)
                for w,v in c.items(): out[prefix+'_'+nm+('_life' if w is None else f'_{w}d')][ii]=v
    return pd.DataFrame(out)
def safe(x): return re.sub(r'[^A-Za-z0-9]+','_',str(x)).strip('_')[:34]
