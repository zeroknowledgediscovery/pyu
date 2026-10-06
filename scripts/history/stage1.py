import pandas as pd, numpy as np, re, json, collections, warnings
from pathlib import Path
from openpyxl import load_workbook
warnings.filterwarnings('ignore')
B=Path('/mnt/data/joel_ua'); OUT=B/'pyuria_longitudinal_full.csv'

def xlsx(path, cols):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; it=ws.iter_rows(values_only=True); hdr=list(next(it)); ids=[hdr.index(c) for c in cols]
    return pd.DataFrame((tuple(r[i] for i in ids) for r in it),columns=cols)
def dt(x): return pd.to_datetime(x,errors='coerce',format='mixed')
def clean_mrn(s): return pd.to_numeric(s,errors='coerce').astype('Int64')
def safe(x): return re.sub(r'[^A-Za-z0-9]+','_',str(x)).strip('_')[:36]

F=pd.read_csv(B/'pyuria_lsm_dataset.csv',low_memory=False); F['UA_DTM']=dt(F.UA_DTM); F['Anon_MRN']=clean_mrn(F.Anon_MRN); F['Anon_Encounter_ID']=clean_mrn(F.Anon_Encounter_ID); F['_row']=np.arange(len(F)); allm=set(F.Anon_MRN.dropna().astype(int))
sp=pd.read_csv('/mnt/data/joel_mc/split_patients_seed20261006.csv'); train=set(pd.to_numeric(sp.loc[sp.split.astype(str).str.lower().eq('train'),'Anon_MRN'],errors='coerce').dropna().astype(int))
left0=F[['_row','Anon_MRN','UA_DTM']].dropna(subset=['Anon_MRN','UA_DTM']).copy(); left0['Anon_MRN']=left0.Anon_MRN.astype(int)

def cum_at(h, qtimes):
    # h columns Anon_MRN,event_dt; qtimes columns _row,Anon_MRN,qtime
    if h.empty: return np.zeros(len(F),float)
    z=h[['Anon_MRN','event_dt']].dropna().copy(); z['Anon_MRN']=z.Anon_MRN.astype(int); z=z[z.Anon_MRN.isin(allm)].sort_values(['event_dt','Anon_MRN']); z['cum']=z.groupby('Anon_MRN').cumcount()+1
    l=qtimes.sort_values(['qtime','Anon_MRN'])
    r=pd.merge_asof(l,z[['Anon_MRN','event_dt','cum']],left_on='qtime',right_on='event_dt',by='Anon_MRN',direction='backward',allow_exact_matches=False)
    out=np.zeros(len(F),float); out[r._row.values.astype(int)]=r['cum'].fillna(0).values; return out

def add_counts(prefix,h,timecol,windows=(7,30,365,None)):
    hh=h[['Anon_MRN',timecol]].rename(columns={timecol:'event_dt'}).copy(); life=cum_at(hh,left0.rename(columns={'UA_DTM':'qtime'}))
    for w in windows:
        if w is None: F[prefix+'_life']=life.astype(np.int32)
        else:
            q=left0.copy(); q['qtime']=q.UA_DTM-pd.to_timedelta(w,unit='D'); past=cum_at(hh,q[['_row','Anon_MRN','qtime']]); F[f'{prefix}_{w}d']=(life-past).clip(min=0).astype(np.int32)

def latest_before(h,timecol,valcol,prefix,cut_hours=0):
    z=h[['Anon_MRN',timecol,valcol]].dropna(subset=['Anon_MRN',timecol]).copy(); z['Anon_MRN']=z.Anon_MRN.astype(int); z=z[z.Anon_MRN.isin(allm)].sort_values([timecol,'Anon_MRN']); l=left0.copy(); l['qtime']=l.UA_DTM-pd.to_timedelta(cut_hours,unit='h'); l=l.sort_values(['qtime','Anon_MRN'])
    r=pd.merge_asof(l,z[['Anon_MRN',timecol,valcol]],left_on='qtime',right_on=timecol,by='Anon_MRN',direction='backward',allow_exact_matches=False)
    vals=np.full(len(F),np.nan); rec=np.full(len(F),np.nan); ii=r._row.values.astype(int); vals[ii]=pd.to_numeric(r[valcol],errors='coerce').values; rec[ii]=(r.qtime-r[timecol]).dt.total_seconds().div(3600).values; F[prefix]=vals; F[prefix+'_rec_h']=rec

print('encounters',flush=True)
enc=xlsx(B/'anonymized_PHI_UA_ENCOUNTER_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','ADMT_DT','ZIP_CD','GENDER_IDENTITY']); enc['Anon_MRN']=clean_mrn(enc.Anon_MRN); enc['Anon_Encounter_ID']=clean_mrn(enc.Anon_Encounter_ID); enc['ADMT_DT']=dt(enc.ADMT_DT); F=F.merge(enc[['Anon_Encounter_ID','ZIP_CD','GENDER_IDENTITY']].drop_duplicates('Anon_Encounter_ID'),on='Anon_Encounter_ID',how='left'); F['ZIP3']=F.ZIP_CD.astype(str).str.replace(r'\.0$','',regex=True).str[:3].where(F.ZIP_CD.notna(),''); add_counts('prior_enc',enc,'ADMT_DT',(7,30,365,None))

print('diagnoses',flush=True)
dx=xlsx(B/'anonymized_PHI_UA_DIAGNOSIS_MRNREFORMAT.xlsx',['Anon_MRN','ADMT_DT','DIAGNOSIS']); dx['Anon_MRN']=clean_mrn(dx.Anon_MRN); dx['ADMT_DT']=dt(dx.ADMT_DT); dx=dx[dx.Anon_MRN.isin(allm)]; dx['icd']=dx.DIAGNOSIS.fillna('').astype(str).str.upper().str.replace('.','',regex=False); add_counts('dx',dx,'ADMT_DT',(30,365,None))
for nm,pat in {'uti':r'^N39','pyelo':r'^N1[0-2]','urinary':r'^R3','fever':r'^R50','abdpain':r'^R10','renal':r'^N[0-2]','diabetes':r'^E1[0-4]'}.items(): add_counts('dx_'+nm,dx[dx.icd.str.match(pat,na=False)],'ADMT_DT',(30,365,None))

print('problem list',flush=True)
pr=xlsx(B/'anonymized_PHI_UA_PROBLEM_LIST_MRNREFORMAT.xlsx',['Anon_MRN','PROBLEM_ICD10_LIST','CHRONIC_YN','NOTED_DATE','DATE_OF_ENTRY']); pr['Anon_MRN']=clean_mrn(pr.Anon_MRN); pr['event_dt']=dt(pr.DATE_OF_ENTRY).fillna(dt(pr.NOTED_DATE)); pr=pr[pr.Anon_MRN.isin(allm)]; add_counts('problem',pr,'event_dt',(365,None)); add_counts('problem_chronic',pr[pr.CHRONIC_YN.astype(str).str.upper().eq('Y')],'event_dt',(365,None))

print('outpatient rx',flush=True)
rx=xlsx(B/'anonymized_PHI_UA_OUTPATIENT_RX_MRNREFORMAT.xlsx',['Anon_MRN','ORDER_START_DATE','ORDER_STOP_DATE','DISPLAY_NAME','DESCRIPTION']); rx['Anon_MRN']=clean_mrn(rx.Anon_MRN); rx['event_dt']=dt(rx.ORDER_START_DATE); rx=rx[rx.Anon_MRN.isin(allm)]; txt=(rx.DISPLAY_NAME.fillna('').astype(str)+' '+rx.DESCRIPTION.fillna('').astype(str)).str.lower(); abxpat=r'amoxic|ampic|penicill|cephal|cef[a-z]|sulfameth|trimeth|nitrofur|ciprof|levof|azith|doxy|clinda|vancom|gentamic|tobram|meropen|piperac|linezolid|metronid'; rx['abx']=txt.str.contains(abxpat,regex=True,na=False); add_counts('rx',rx,'event_dt',(30,365,None)); add_counts('rx_abx',rx[rx.abx],'event_dt',(30,365,None))

print('imaging',flush=True)
im=xlsx(B/'anonymized_PHI_UA_IMAGING_EPIC_MRNREFORMAT.xlsx',['Anon_MRN','FINAL_DT_TM','PROCDR1_NAME']); im['Anon_MRN']=clean_mrn(im.Anon_MRN); im['event_dt']=dt(im.FINAL_DT_TM); im=im[im.Anon_MRN.isin(allm)]; add_counts('imaging',im,'event_dt',(30,365,None))
for nm,pat in {'us':r'^US','ct':r'^CT','xr':r'^XR','mri':r'^(MRI|MR )'}.items(): add_counts('img_'+nm,im[im.PROCDR1_NAME.fillna('').astype(str).str.upper().str.match(pat,na=False)],'event_dt',(365,None))

F=F.drop(columns=['_row'],errors='ignore').replace([np.inf,-np.inf],np.nan)
F.to_csv(B/'pyuria_longitudinal_stage1.csv',index=False)
print('STAGE1 DONE',F.shape,flush=True)
