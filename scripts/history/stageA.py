import pandas as pd, numpy as np, re, json, collections, warnings
from pathlib import Path
from openpyxl import load_workbook
warnings.filterwarnings('ignore')
B=Path('/mnt/data/joel_ua'); OUT=B/'pyuria_longitudinal_full.csv'

def xlsx(path, cols):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; it=ws.iter_rows(values_only=True); hdr=list(next(it)); ids=[hdr.index(c) for c in cols]
    df=pd.DataFrame((tuple(r[i] for i in ids) for r in it),columns=cols); wb.close(); return df
def dt(x): return pd.to_datetime(x,errors='coerce',format='mixed')
def mrn(x): return pd.to_numeric(x,errors='coerce').astype('Int64')
def safe(x): return re.sub(r'[^A-Za-z0-9]+','_',str(x)).strip('_')[:34]
F=pd.read_csv(B/'pyuria_lsm_dataset.csv',low_memory=False); F['UA_DTM']=dt(F.UA_DTM); F['Anon_MRN']=mrn(F.Anon_MRN); F['Anon_Encounter_ID']=mrn(F.Anon_Encounter_ID); allm=set(F.Anon_MRN.dropna().astype(int))
sp=pd.read_csv('/mnt/data/joel_mc/split_patients_seed20261006.csv'); train=set(pd.to_numeric(sp.loc[sp.split.astype(str).str.lower().eq('train'),'Anon_MRN'],errors='coerce').dropna().astype(int))
ua={int(m):(np.asarray(idx,dtype=int),F.loc[idx,'UA_DTM'].values.astype('datetime64[ns]')) for m,idx in F.groupby('Anon_MRN').groups.items() if pd.notna(m)}

def count_windows(t,uts,windows):
    t=np.sort(np.asarray(t,dtype='datetime64[ns]')); r=np.searchsorted(t,uts,'left'); out={}
    for w in windows:
        out[w]=r if w is None else r-np.searchsorted(t,uts-np.timedelta64(int(w),'D'),'left')
    return out

def add_source_counts(h,timecol,prefix,windows=(30,365,None),mask_defs=None):
    h=h[h.Anon_MRN.notna() & h[timecol].notna()].copy(); h.Anon_MRN=h.Anon_MRN.astype(int); h=h[h.Anon_MRN.isin(allm)]
    def fill(z,nm):
        arrays={w:np.zeros(len(F),np.int32) for w in windows}
        for m,g in z.groupby('Anon_MRN',sort=False):
            if m not in ua: continue
            ii,uts=ua[m]; c=count_windows(g[timecol].values.astype('datetime64[ns]'),uts,windows)
            for w,v in c.items(): arrays[w][ii]=v
        for w,v in arrays.items(): F[nm+('_life' if w is None else f'_{w}d')]=v
    fill(h,prefix)
    if mask_defs:
        for nm,func in mask_defs.items():
            mask=np.asarray(func(h),dtype=bool)
            fill(h.loc[mask],prefix+'_'+nm)

print('enc',flush=True)
enc=xlsx(B/'anonymized_PHI_UA_ENCOUNTER_MRNREFORMAT.xlsx',['Anon_Encounter_ID','Anon_MRN','ADMT_DT','ZIP_CD','GENDER_IDENTITY']); enc['Anon_MRN']=mrn(enc.Anon_MRN); enc['Anon_Encounter_ID']=mrn(enc.Anon_Encounter_ID); enc['ADMT_DT']=dt(enc.ADMT_DT); F=F.merge(enc[['Anon_Encounter_ID','ZIP_CD','GENDER_IDENTITY']].drop_duplicates('Anon_Encounter_ID'),on='Anon_Encounter_ID',how='left'); F['ZIP3']=F.ZIP_CD.astype(str).str.replace(r'\.0$','',regex=True).str[:3].where(F.ZIP_CD.notna(),''); add_source_counts(enc,'ADMT_DT','prior_enc',(7,30,365,None))

print('dx',flush=True)
dx=xlsx(B/'anonymized_PHI_UA_DIAGNOSIS_MRNREFORMAT.xlsx',['Anon_MRN','ADMT_DT','DIAGNOSIS']); dx['Anon_MRN']=mrn(dx.Anon_MRN); dx['ADMT_DT']=dt(dx.ADMT_DT); dx['icd']=dx.DIAGNOSIS.fillna('').astype(str).str.upper().str.replace('.','',regex=False)
dx['uti']=dx.icd.str.match(r'^N39',na=False); dx['pyelo']=dx.icd.str.match(r'^N1[0-2]',na=False); dx['urinary']=dx.icd.str.match(r'^R3',na=False); dx['renal']=dx.icd.str.match(r'^N[0-2]',na=False)
add_source_counts(dx,'ADMT_DT','dx',(30,365,None),{'uti':lambda g:g.uti,'pyelo':lambda g:g.pyelo,'urinary':lambda g:g.urinary,'renal':lambda g:g.renal}); del dx

print('problem',flush=True)
pr=xlsx(B/'anonymized_PHI_UA_PROBLEM_LIST_MRNREFORMAT.xlsx',['Anon_MRN','CHRONIC_YN','NOTED_DATE','DATE_OF_ENTRY']); pr['Anon_MRN']=mrn(pr.Anon_MRN); pr['event_dt']=dt(pr.DATE_OF_ENTRY).fillna(dt(pr.NOTED_DATE)); add_source_counts(pr,'event_dt','problem',(365,None),{'chronic':lambda g:g.CHRONIC_YN.astype(str).str.upper().eq('Y')}); del pr

print('rx',flush=True)
rx=xlsx(B/'anonymized_PHI_UA_OUTPATIENT_RX_MRNREFORMAT.xlsx',['Anon_MRN','ORDER_START_DATE','DISPLAY_NAME','DESCRIPTION']); rx['Anon_MRN']=mrn(rx.Anon_MRN); rx['event_dt']=dt(rx.ORDER_START_DATE); txt=(rx.DISPLAY_NAME.fillna('').astype(str)+' '+rx.DESCRIPTION.fillna('').astype(str)).str.lower(); abxpat=r'amoxic|ampic|penicill|cephal|cef[a-z]|sulfameth|trimeth|nitrofur|ciprof|levof|azith|doxy|clinda|vancom|gentamic|tobram|meropen|piperac|linezolid|metronid'; rx['abx']=txt.str.contains(abxpat,regex=True,na=False); add_source_counts(rx,'event_dt','rx',(30,365,None),{'abx':lambda g:g.abx}); del rx
F.to_csv(B/'stageA.csv',index=False); print('STAGEA DONE',F.shape,flush=True)
